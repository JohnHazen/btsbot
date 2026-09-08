from __future__ import annotations

import log

import datetime
from dataclasses import dataclass,field
from typing import *  # type:ignore
from sqlalchemy import ForeignKey
from sqlalchemy import Table
from sqlalchemy import Column
from sqlalchemy import Integer, String, DateTime
from sqlalchemy import create_engine
from sqlalchemy import select
from sqlalchemy.sql import func
from sqlalchemy.orm import DeclarativeBase
from sqlalchemy.orm import Mapped
from sqlalchemy.orm import MappedAsDataclass
from sqlalchemy.orm import mapped_column
from sqlalchemy.orm import relationship
from sqlalchemy.orm import sessionmaker
#from sqlalchemy.orm import Session
#from . import Session

import itertools
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Alignment
from openpyxl.styles.borders import Border, Side
from openpyxl.utils import get_column_letter

import utility
from utility import select_quals,assign_parts
#from utility import select_quals,DAYS_TO_QUAL_EXPIRATION,WARN_DAYS_PRIOR_TO_QUAL_EXPIRATION

class Base(MappedAsDataclass, DeclarativeBase):
    pass

# association tables are declared without ORM 

song_tag_table = Table("SongTag",Base.metadata,
    Column("song_id", ForeignKey("Song.id"), primary_key=True),
    Column("tag_id", ForeignKey("Tag.id"), primary_key=True),
    )

gig_singer_table = Table("GigSinger",Base.metadata,
    Column("user_id", ForeignKey("User.id"), primary_key=True),
    Column("gig_id", ForeignKey("Gig.id"), primary_key=True),
    )


class User(Base):
    __tablename__ = "User"

    id: Mapped[int] = mapped_column(init=False, primary_key=True)
    name: Mapped[str] = mapped_column(String(40))
    slack_id: Mapped[str] = mapped_column(String(40), default='')
    auths: Mapped[str] = mapped_column(String(255), default='')
    org: Mapped[str] = mapped_column(String(40), default='BTS')
    override: Mapped[str] = mapped_column(String(40), default='Normal')
    utility: Mapped[int] = mapped_column(Integer, default = 0)

    def __hash__(self):
        return hash((self.name,self.id))

    def __lt__(self,other):
        return self.utility < other.utility

    def __eq__(self,other):
        return self.id == other.id

    def set_slack_id(self,slack_id):
        self.slack_id = slack_id
        session.add(self)
        session.commit()

    def add_auth(self,auth):
        auth_list = self.auths.split(',')
        if auth in auth_list:
            return
        auth_list.append(auth)
        if "" in auth_list:
            auth_list.remove("")
        self.auths = ','.join(auth_list)
        session.add(self)
        session.commit()

    def update(self,name=None,auth_list=None):
        log.debug(f"User update - name={name}  auth_list={auth_list}")
        if name is not None:
            self.name = name
        if auth_list is not None:
            self.auths = ",".join(auth_list)
        session.add(self)
        session.commit()

    def authorized(self,auth):
        auth_list = self.auths.split(',')
        if auth in auth_list:
            return True
        return False

class Song(Base):
    __tablename__ = "Song"

    id: Mapped[int] = mapped_column(init=False, primary_key=True)
    name: Mapped[str] = mapped_column(String(40))
    tags: Mapped[List[Tag]] = relationship(
        secondary=song_tag_table, back_populates="songs"
        )
    parts: Mapped[List[SongVoicePart]] = relationship()
    difficulty: Mapped[int] = mapped_column(Integer, default = 0)
    duration: Mapped[int] = mapped_column(Integer, init=False, nullable=True) #in seconds
    written_key: Mapped[str] = mapped_column(String(8),init=False, nullable=True)
    performance_key: Mapped[str] = mapped_column(String(8),init=False, nullable=True)

    def __lt__(self,other):
        return self.name < other.name

    def update(self,name=None,tags=None,difficulty=None,written_key=None,performance_key=None,duration=None):
        if name is not None:
            self.name = name
        if difficulty is not None:
            self.difficulty = difficulty
        if written_key is not None:
            self.written_key = written_key
        if performance_key is not None:
            self.performance_key = performance_key
        if duration is not None:
            self.duration = duration
        if tags is not None:
            target_tags = set(tags)
            current_tags = set(self.tags)
            diff_tags = target_tags ^ current_tags
            add_tags = target_tags - current_tags
            del_tags = current_tags - target_tags
            for tag in diff_tags:
                session.add(tag)
            # may need to actually add/delete song in tags...
            self.tags = tags
        session.add(self)
        session.commit()

class Tag(Base):
    __tablename__ = "Tag"

    id: Mapped[int] = mapped_column(init=False, primary_key=True)
    name: Mapped[str] = mapped_column(String(40))
    songs: Mapped[List[Song]] = relationship(
        secondary=song_tag_table, back_populates="tags"
        )
    sort_order: Mapped[int] = mapped_column(Integer, default = 0)

    def __hash__(self):
        return self.id

    def __lt__(self,other):
        if self.sort_order != other.sort_order:
            return self.sort_order < other.sort_order
        return self.name < other.name

class VoicePart(Base):
    __tablename__ = "VoicePart"

    id: Mapped[int] = mapped_column(init=False, primary_key=True)
    name: Mapped[str] = mapped_column(String(40))
    range_low: Mapped[int] = mapped_column(init=False, nullable=True)
    range_high: Mapped[int] = mapped_column(init=False, nullable=True)

class SongVoicePart(Base):
    __tablename__ = "SongVoicePart"

    # TODO fix below to be ints (I think just annotations, so shouldn't break anything
    song_id: Mapped[str] = mapped_column(ForeignKey("Song.id"), init=False, primary_key=True)
    part_id: Mapped[str] = mapped_column(ForeignKey("VoicePart.id"), init=False, primary_key=True)
    part: Mapped["VoicePart"] = relationship(init=False)
    range_low: Mapped[int] = mapped_column(init=False, nullable=True)
    range_high: Mapped[int] = mapped_column(init=False, nullable=True)
    sort_order: Mapped[int] = mapped_column(init=False, nullable=True)

    def __lt__(self,other):
        if self.sort_order == other.sort_order:
            return self.part_id < other.part_id
        return self.sort_order < other.sort_order
        

    def update(self,sort_order=None,range_low=None,range_high=None):
        log.debug(f"SongVoicePart update - sort_order={sort_order} range_low={range_low} range_high={range_high}")
        if sort_order is not None:
            self.sort_order = sort_order
        if range_low is not None:
            self.range_low = range_low
        if range_high is not None:
            self.range_high = range_high
        session.add(self)
        session.commit()

class SongVoicePartQual(Base):
    __tablename__ = "SongVoicePartQual"

    id: Mapped[int] = mapped_column(init=False, primary_key=True)
    # TODO fix below to be ints (I think just annotations, so shouldn't break anything
    user_id: Mapped[str] = mapped_column(ForeignKey("User.id"), init=False)
    song_id: Mapped[str] = mapped_column(ForeignKey("Song.id"), init=False)
    part_id: Mapped[str] = mapped_column(ForeignKey("VoicePart.id"), init=False)
    date_time: Mapped[datetime.datetime] = mapped_column(DateTime(timezone=True), init=False, server_default=func.now())
    performance_ready: Mapped[bool] = mapped_column(init=False, unique=False, default=True)
    comments: Mapped[str] = mapped_column(String(255),init=False, nullable=True)

    def __lt__(self,other):
        return self.id < other.id

class GigSetListItem(Base):
    ''' use these items to build a setlist.  They may refer to a song, or just be used for the comment
        to indicate a speaking part...
    '''
    __tablename__ = "GigSetListItem"

    id: Mapped[int] = mapped_column(init=False, primary_key=True)
    gig_id: Mapped[int] = mapped_column(ForeignKey("Gig.id"), init=False)
    song_id: Mapped[int] = mapped_column(ForeignKey("Song.id"), init=False, nullable=True)
    sort_order: Mapped[int] = mapped_column(init=False, nullable=True)
    duration: Mapped[int] = mapped_column(init=False, nullable=True) # seconds
    song_key: Mapped[str] = mapped_column(String(8),init=False, nullable=True)
    in_setlist: Mapped[bool] = mapped_column(init=False, unique=False, default=False)
    comments: Mapped[str] = mapped_column(String(255),init=False, nullable=True)

    def __lt__(self,other):
        if self.sort_order == other.sort_order and self.in_setlist != other.in_setlist:
            return self.in_setlist
        return self.sort_order < other.sort_order

    def update(self,sort_order=None,in_setlist=None,song_key=None,duration=None,comments=None):
        log.debug(f"SetListItem update - sort_order={sort_order}  in_setlist={in_setlist} song_key={song_key} duration={duration},comments={comments}")
        if sort_order is not None:
            self.sort_order = sort_order
        if in_setlist is not None:
            self.in_setlist = in_setlist
        if song_key is not None:
            self.song_key = song_key
        if comments is not None:
            self.comments = comments
        if duration is not None:
            self.duration = duration
        session.add(self)
        session.commit()

class GigSongVoicePartAssignment(Base):
    __tablename__ = "GigSongVoicePartAssignment"

    id: Mapped[int] = mapped_column(init=False, primary_key=True)
    gig_id: Mapped[int] = mapped_column(ForeignKey("Gig.id"), init=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("User.id"), init=True)
    song_id: Mapped[int] = mapped_column(ForeignKey("Song.id"), init=True)
    part_id: Mapped[int] = mapped_column(ForeignKey("VoicePart.id"), init=True)
    date_time: Mapped[datetime.datetime] = mapped_column(DateTime(timezone=True), init=False, nullable=True)
    performance_ready: Mapped[bool] = mapped_column(init=False, unique=False, default=True)
    waiver: Mapped[bool] = mapped_column(init=False, unique=False, default=False)
    sort_order: Mapped[int] = mapped_column(init=False, nullable=True)
    comments: Mapped[str] = mapped_column(String(255),init=False, nullable=True)

    def emoji(self):
        if self.waiver:
            return ":white_check_mark:"
        if self.performance_ready:
            return ":large_green_circle:"
        if self.date_time is not None:
            return ":red_circle:"
        return ":negative_squared_cross_mark:"
        return ":black_circle:"
        return ":x:"
        return ":large_yellow_circle:"
        return ":thumbsup:"


    def __lt__(self,other):
        return self.sort_order < other.sort_order
        return self.id < other.id

    def update(self,sort_order=None,performance_ready=None,waiver=None,comments=None):
        log.debug(f"PartAssignment update - sort_order={sort_order}  performance_ready={performance_ready} waiver={waiver} comments={comments}")
        if sort_order is not None:
            self.sort_order = sort_order
        if performance_ready is not None:
            self.performance_ready = performance_ready
        if waiver is not None:
            self.waiver = waiver
        if comments is not None:
            self.comments = comments
        session.add(self)
        session.commit()

class Gig(Base):
    __tablename__ = "Gig"

    id: Mapped[int] = mapped_column(init=False, primary_key=True)
    location: Mapped[str] = mapped_column(String(255))
    description: Mapped[str] = mapped_column(String(255))
    comments: Mapped[str] = mapped_column(String(255), init=False, default="")
    singers: Mapped[List[User]] = relationship(
        secondary=gig_singer_table)
    setlist: Mapped[List[GigSetListItem]] = relationship(init=False)
    assignments: Mapped[List[GigSongVoicePartAssignment]] = relationship(init=False)
    date_time: Mapped[datetime.datetime] = mapped_column(DateTime(timezone=True), init=False, nullable=True)
    duration: Mapped[int] = mapped_column(Integer, init=False, default=0)   # in minutes
    active: Mapped[bool] = mapped_column(init=False, unique=False, default=True)
    thread_ts: Mapped[str] = mapped_column(String(255),init=False,default="")

    def __lt__(self,other):
        return self.date_time < other.date_time

    def update(self,location=None,description=None,comments=None,singers=None,
                date_time=None,duration=None,active=None):
        log.debug(f"Gig update - location={location} description={description} comments={comments} singers={singers} date_time={date_time} duration={duration} active={active}")
        if location is not None:
            self.location = location
        if description is not None:
            self.description = description
        if comments is not None:
            self.comments = comments
        if date_time is not None:
            self.date_time = date_time
        if duration is not None:
            self.duration = duration
        if active is not None:
            self.active = active
        session.add(self)
        session.commit()
        if singers is not None:
            self.update_singers(singers)

    def setlist_duration(self):
        duration = 0
        for item in sorted(self.setlist):
            if item.comments == "-- END SETLIST --":
                return duration
            try:
                duration += item.duration
            except:
                log.warning(f"error computing set duration for gig {self.description}")
        log.error(f"error computing set duration for gig {self.description}.  No END marker.")

    def delete_setlist_item(self,item):
        '''
            remove setlist_item
            - then remove all relevant assignments
        '''
        log.debug(f"Gig delete_setlist_item - item={item}")
        if item not in self.setlist:
            log.error(f"tried to delete item not in this gig's setlist.")
            return
        if item.song_id is not None:
            assignments_left_over = []
            for a in self.assignments:
                if a.song_id == item.song_id:
                    session.delete(a)
                else:
                    assignments_left_over.append(a)
        self.setlist.remove(item)
        session.delete(item)
        self.assignments = assignments_left_over
        session.add(self)
        session.commit()

    def add_setlist_item(self,song_id=None,comments=None):
        '''
            add setlist_item
            - then make sure to add singers' assignments
        '''
        log.debug(f"Gig add_setlist_item - song_id={song_id} comments={comments}")
        end_setlist_item = [i for i in self.setlist if i.comments == '-- END SETLIST --'][0]
        insert_point = self.setlist.index(end_setlist_item)
        item=GigSetListItem()
        item.gig_id = self.id
        item.comments = comments
        item.sort_order = end_setlist_item.sort_order
        item.in_setlist = True
        if song_id is not None:
            item.song_id = song_id

            # add new assignments for new singers/songs
            singer_quals = select_quals(quals,song_id=song_id,
                    user_ids=[x.id for x in self.singers],only_most_recent=True)
            for qual in singer_quals:
                expiry = QualExpirationData(qual,target_date=self.date_time)
                part_assignment = GigSongVoicePartAssignment(
                            gig_id=self.id,
                            user_id=qual.user_id,
                            song_id=qual.song_id,
                            part_id=qual.part_id)
                part_assignment.date_time=qual.date_time
                part_assignment.performance_ready = expiry.valid
                part_assignment.sort_order = expiry.sort_order
                if not expiry.valid:
                    part_assignment.sort_order += 1000
                self.assignments.append(part_assignment)

            self.assignments = sorted(self.assignments)

            # now add non-qual assignments for dropdown menus
            song = songs_by_id[item.song_id]
            item.duration = song.duration
            item.song_key = song.performance_key if song.performance_key else song.written_key
            for singer in self.singers:
                for part in sorted(song.parts):
                    if not [a for a in self.assignments if (a.part_id == part.part.id 
                            and a.song_id == item.song_id and a.user_id == singer.id)]:
                        # no existing qual-based-assignment, so create one
                        part_assignment = GigSongVoicePartAssignment(
                                    gig_id=self.id,
                                    user_id=singer.id,
                                    song_id=song.id,
                                    part_id=part.part.id)
                        part_assignment.performance_ready = False
                        part_assignment.sort_order = 2000
                        self.assignments.append(part_assignment)

            log.debug(f"assigning parts for {song.name}")
            assign_parts(song,[a for a in self.assignments if a.song_id==song.id])


        self.setlist.insert(insert_point,item)
        session.add(item)
        session.add(self)
        session.commit()

    def update_singers(self,singers=None):
        ''' add or subtract singers
                pass end_state singer list
            Steps:
            - add/remove actual singers
            - remove (now) unused assignments
            - add any new songs to setlist
            - add new assignments for new singers/songs
            - auto-assign new songs
        '''
        log.debug(f"Gig update_singers - singers={singers}")
        if singers is None:
            return
        gig_singers = set(self.singers)
        new_singers = set(singers)
        if gig_singers == new_singers:
            log.debug(f"Gig update_singers - Singers didn't change.  Nothing to do.")
            return
        # add/remove actual singers
        self.singers = singers
        added_singers = new_singers - gig_singers
        deleted_singers = gig_singers - new_singers
        log.debug(f"adding singers: {added_singers}  Removing singers: {deleted_singers}")
        added_singer_ids = [x.id for x in added_singers]
        deleted_singer_ids = [x.id for x in deleted_singers]

        # remove (now) unused assignments
        kept_assignments = [x for x in self.assignments if x.user_id not in deleted_singer_ids]
        for assignment in [x for x in self.assignments if x.user_id in deleted_singer_ids]:
            session.delete(assignment)
        self.assignments = kept_assignments

        # add any new songs to setlist
        self.setlist = sorted(self.setlist)
        singer_quals = select_quals(quals,user_ids=[x.id for x in singers],only_most_recent=True)
        set_song_ids = [x.song_id for x in self.setlist]
        if self.setlist:
            end_setlist_item = [i for i in self.setlist if i.comments == '-- END SETLIST --'][0]
            insert_point = self.setlist.index(end_setlist_item)
            last_setlist_item_index = insert_point - 1
            if last_setlist_item_index < 0:
                last_setlist_item_index = 0
            sort_order = self.setlist[last_setlist_item_index].sort_order
        else:
            end_setlist_item = GigSetListItem()
            end_setlist_item.sort_order = 50
            end_setlist_item.comments = "-- END SETLIST --"
            end_setlist_item.in_setlist = False
            self.setlist.append(end_setlist_item)
            insert_point = 0
            sort_order = 0
        sorted_songs = sorted([(utility.qual_strength(singer_quals,s),s) for s in songs],reverse=True)
        performable = True
        new_setlist_items = []
        for ((qual_strength,missing_parts),song) in sorted_songs:
            if song.id in set_song_ids:
                log.debug(f"skipping {song.name} -- already in setlist")
                continue
            sort_order += 1
            if qual_strength < 0 and performable:
                performable = False
                insert_point += 1 # move from inserting *before* end-marker, to inserting after
                end_setlist_item.sort_order = sort_order
                sort_order += 100
                log.debug(f"first non-performable: {song.name}")
            if qual_strength < -1:
                break
            log.debug(f"adding song: {song.name}")
            setlist_item=GigSetListItem()
            setlist_item.gig_id=self.id
            setlist_item.song_id=song.id
            setlist_item.duration=song.duration
            setlist_item.song_key = song.performance_key if song.performance_key else song.written_key
            setlist_item.sort_order=sort_order
            setlist_item.in_setlist=performable
            self.setlist.insert(insert_point,setlist_item)
            new_setlist_items.append(setlist_item)
            insert_point += 1

        # add new assignments for new singers/songs
        set_song_ids = [x.song_id for x in self.setlist]
        singer_quals = select_quals(quals,user_ids=[x.id for x in self.singers],only_most_recent=True)
        for qual in singer_quals:
            if qual.song_id not in set_song_ids:
                continue
            # check if there's already an assignment for this qual
            qual_assignments = [x for x in self.assignments if (x.user_id==qual.user_id 
                and x.song_id == qual.song_id and x.part_id==qual.part_id)]
            if qual_assignments:
                # leave sort_order, in case it's been assigned, but update date_time and performance_ready
                qual_assignments[0].date_time = qual.date_time
                qual_assignments[0].performance_ready = qual.performance_ready
            else:
                #create new assignment
                expiry = QualExpirationData(qual,target_date=self.date_time)
                part_assignment = GigSongVoicePartAssignment(
                            gig_id=self.id,
                            user_id=qual.user_id,
                            song_id=qual.song_id,
                            part_id=qual.part_id)
                part_assignment.date_time=qual.date_time
                part_assignment.performance_ready = expiry.valid
                part_assignment.sort_order = expiry.sort_order
                if not expiry.valid:
                    part_assignment.sort_order += 1000
                self.assignments.append(part_assignment)
        self.assignments = sorted(self.assignments)

        # now add non-qual assignments so singer shows up in pulldowns for non-qual gigs.
        # and auto-assign parts for new setlist_items and older items not in setlist 
        #   (in case adding new singers makes a quartet)  Figure items in setlist have already 
        #   been assigned, so leave them alone.
        in_setlist = True
        for setlist_item in self.setlist:
            if setlist_item is end_setlist_item:
                in_setlist = False
            if setlist_item.song_id is None:
                continue
            song = songs_by_id[setlist_item.song_id]
            #for singer in added_singers:
            for singer in self.singers:
                for part in sorted(song.parts):
                    if not [a for a in self.assignments if (a.part_id == part.part.id 
                            and a.song_id == setlist_item.song_id and a.user_id == singer.id)]:
                        # no existing qual-based-assignment, so create one
                        part_assignment = GigSongVoicePartAssignment(
                                    gig_id=self.id,
                                    user_id=singer.id,
                                    song_id=song.id,
                                    part_id=part.part.id)
                        part_assignment.performance_ready = False
                        part_assignment.sort_order = 2000
                        self.assignments.append(part_assignment)

            # auto-assign new songs
            if setlist_item in new_setlist_items or not in_setlist:
                log.debug(f"assigning parts for {song.name}")
                assign_parts(song,[a for a in self.assignments if a.song_id==song.id])

        session.add(self)
        session.commit()

class Note(object):
    names = ["C","Db","D","Eb","E","F","Gb","G","Ab","A","Bb","B"]

    def __init__(self,name=None,value=None):
        if name is None and value is None:
            raise ValueError("name or value must be specified")
        if value is not None:
            self.value = value
            octave = int(value / 12)
            note = value % 12
            self.name = f"{self.names[note]}{octave}"
        if name is not None:
            self.name = name
            octave = int(name[-1:])
            note_value = self.names.index(name[:-1])
            try:
                octave = int(name[-1:])
            except:
                raise ValueError(f"note name must be one of {self.names} followed by an octave number.  Got {name}")
            try:
                note_value = self.names.index(name[:-1])
            except:
                raise ValueError(f"note name must be one of {self.names} followed by an octave number.  Got {name}")
            self.value = 12*octave + note_value

class PersistentKeyValue(Base):
    __tablename__ = "KeyValue"
    # TODO if we move to a real database, add sanity checks for length.  sqlite doesn't actually care

    id: Mapped[int] = mapped_column(init=False, primary_key=True)
    key: Mapped[str] = mapped_column(String(255))
    value: Mapped[str] = mapped_column(String(255))

    def __hash__(self):
        return self.key


@dataclass
class Quartet:
    """
    contains 4 Singer objects, plus various scoring methods and
    maybe containers for quartets related to this one.
    """
    name: str
    singers: set(User)
    score: float

    def __lt__(self,other):
        return self.score < other.score

    def __repr__(self):
        return f"{self.score}  {' / '.join([x.name for x in sorted(self.singers)])}"

@dataclass
class QuartetSet:
    """
    contains N Quartet objects, plus various scores 
    """
    quartets: set(Quartet)
    name: str = ""
    min: float | None = 0
    avg: float = 0
    len: int = 0

    def calculate(self):
        """
        compute min/avg/score/len
        """
        self.len = len(self.quartets)
        sum = 0
        self.min = None
        for q in self.quartets:
            sum += q.score
            if self.min is None or self.min > q.score:
                self.min = q.score
        self.avg = sum/self.len if self.len else 0
        if self.min is None:
            self.min = 0

class QualExpirationData:
    '''
    ## Attributes
    days : days until expriation
    status : status text
    emoji : status emoji (red for expired, yellow for expiring soon, green for valid)
    qual_or_date: InitVar
    qual: SongVoicePart = field(init=False)
    date: datetime.date = field(init=False)
    status: str = field(init=False)
    days: int = field(init=False)
    emoji: str = field(init=False)
    '''

    def __init__(self,qual_or_date,target_date=None):
        if target_date is None:
            target_date = datetime.date.today()
            #target_date = datetime.datetime.now()
        else:
            #try to turn datetimes into dates
            try:
                target_date = target_date.date()
            except AttributeError:
                pass
        if type(qual_or_date) == SongVoicePartQual:
            self.qual = qual_or_date
            #self.date = qual_or_date.date_time.date()
            self.date = qual_or_date.date_time
        else:
            self.qual = None
            #self.date = qual_or_date.date()
            self.date = qual_or_date
        # below is to tolerated both date and datetime objects
        try:
            self.date = self.date.date()
        except AttributeError:
            pass
        days_since_qual = target_date - self.date
        self.days = (datetime.timedelta(days=utility.DAYS_TO_QUAL_EXPIRATION) - days_since_qual).days
        if self.days < 0:
            self.emoji = ":red_circle:"
            self.status = f'Expired {-1*self.days} days ago'
            self.valid = False
        elif self.days < utility.WARN_DAYS_PRIOR_TO_QUAL_EXPIRATION:
            self.emoji = ":large_yellow_circle:"
            self.status = f'Expiring in {self.days} days'
            self.valid = True
        else:
            self.emoji = ":large_green_circle:"
            self.status = f'Good for {self.days} more days'
            self.valid = True
        self.sort_order = days_since_qual.days

    def __lt__(self,other):
        return self.days < other.days


def insert_records(records,singers=None,songs=None,parts=None,tags=None):
    _singers = {}
    _songs = {}
    _parts = {}
    _tags = {}

    if singers is not None:
        for s in singers:
            _singers[s.name] = s
    if songs is not None:
        for s in songs:
            _songs[s.name] = s
    if parts is not None:
        for s in parts:
            _parts[s.name] = s
    if tags is not None:
        for s in tags:
            _tags[s.name] = s

    singers = _singers
    songs = _songs
    parts = _parts
    tags = _tags

    #once through first so songs added have all 4 parts.
    for (tag_name,singer_name,org,song_name,part_name,date) in records:
        if part_name is None:
            continue
        try:
            base_part = parts[part_name]
        except KeyError:
            log.info(f"creating (base) part {part_name}")
            base_part = VoicePart(name=part_name)
            parts[part_name] = base_part
            session.add(base_part)
            session.commit()

    for (tag_name, singer_name, org, song_name, part_name, date) in records:
        if None in [tag_name,song_name]:
            log.slack.error(f"Insert Record:  Error.  Song/tag is None.  {tag_name},{singer_name},{org},{song_name},{part_name},{date}")
            continue
        song_only = None in [singer_name, org, part_name, date]
        if song_only:
            log.debug(f"adding only song for {tag_name} - {song_name}")
        try:
            tag = tags[tag_name]
        except KeyError:
            log.info(f"creating tag {tag_name}")
            tag = Tag(name=tag_name,songs=[])
            tags[tag_name] = tag
            session.add(tag)
            session.commit()
        try:
            song = songs[song_name]
        except KeyError:
            log.info(f"creating song {song_name}")
            song = Song(name=song_name, tags=[tag], parts=[])
            songs[song_name] = song
            session.add(song)
            session.commit()
            song_parts = []
            for p_name, part in parts.items():
                song_part = SongVoicePart()
                song_part.part = part
                song_parts.append(song_part)
            song.parts.extend(song_parts)
            session.add_all(song_parts)
            session.add(song)
            session.commit()
        if tag not in song.tags:
            song.tags.append(tag)
            session.add(song)
            session.commit()
        if song_only:
            continue

        try:
            singer = singers[singer_name]
        except KeyError:
            log.info(f"creating singer {singer_name}")
            override = "Normal"
            singer = User(name=singer_name,override=override,org=org)
            singers[singer_name] = singer
            session.add(singer)
            session.commit()
        
        part = [p for p in song.parts if p.part.name == part_name][0]
        log.info(f"creating qual record")
        qual = SongVoicePartQual()
        qual.part_id = part.part.id
        qual.song_id = song.id
        qual.user_id = singer.id
        if type(date) is datetime.datetime:
            qual.date_time = date
        else:
            try:
                qual.date_time = datetime.datetime.strptime(date,'%m/%d/%y*')
            except ValueError:
                log.warning(f"*** date error in date '{date}' of type{type(date)}.  skipping")
                continue
        session.add(qual)
        # add one to utility for each part learned
        singer.utility += 1
        session.add(singer)
        session.commit()

def create_user(name,slack_id=None):
    user = User(name=name,slack_id=slack_id)
    session.add(user)
    session.commit()
    users.append(user)
    users_by_id[user.id] = user
    if slack_id is not None:
        user_id_from_slack_id[slack_id] = user.id
    return user

def create_song(name,tags=[]):
    log.info(f"creating song {name}")
    song = Song(name=name, tags=tags, parts=[])
    session.add(song)
    session.commit()
    song_parts = []
    for part in parts:
        song_part = SongVoicePart()
        song_part.part = part
        song_parts.append(song_part)
    song.parts.extend(song_parts)
    session.add_all(song_parts)
    session.add(song)
    session.commit()
    songs.append(song)
    songs_by_id[song.id] = song
    return song

def delete_qual(qual):
    #first delete from local cache
    log.debug(f"Deleting: {qual}")
    quals.remove(qual)
    #now delete from DB
    session.delete(qual)
    session.commit()

def delete_tag(tag):
    #TODO check if songs associated with the tag before deleting (already doing this elsewhere)
    #first delete from local cache
    tags.remove(tag)
    #now delete from DB
    session.delete(tag)
    session.commit()


def create_tag(name):
    ''' 
        create new song tag with name.  No associated songs yet
    '''
    tag = Tag(name,[])
    session.add(tag)
    session.commit()
    tags.append(tag)
    return tag

def rename_tag(tag_id,name):
    ''' 
        rename song tag with tag_id to name
    '''
    tag = [t for t in tags if t.id==tag_id][0]
    tag.name = name
    session.add(tag)
    session.commit()
    return tag

def create_qual(singer,song,part_name,date=None):
    ''' singer, song are database objects.  part_name is a string matching the part name.
        date is a datetime.date, or None.  If not specified, use today's date.
    '''
    part = [p for p in song.parts if p.part.name.lower() == part_name.lower()][0]
    qual = SongVoicePartQual()
    qual.part_id = part.part.id
    qual.song_id = song.id
    qual.user_id = singer.id
    if date is not None:
        log.debug("qual date is {date}")
        qual.date_time = date
    else:
        qual.date_time = datetime.datetime.today()
    session.add(qual)
    session.commit()
    quals.append(qual)
    return qual

def create_gig(singer_ids,date_time=None,location="",description="",thread_ts=""):
    if date_time is None:
        date_time = datetime.datetime.now() + datetime.timedelta(days=14)
    #gig = Gig(location,description,[users_by_id[s] for s in singer_ids])
    gig = Gig(location,description,[])
    #commit gig to get ID
    session.add(gig)
    session.commit()
    gig.date_time=date_time
    gig.thread_ts=thread_ts
    if description == "" and location == "":
        gig.description = f"Gig ID {gig.id}"
    session.add(gig)
    session.commit()
    gigs.append(gig)

    singers = sorted([users_by_id[x] for x in singer_ids])
    gig.update_singers(singers=singers)
    return gig

    #TODO delete all below, as that's now handled by update_singers


    # set up setlist and possible assignments
    singer_quals = select_quals(quals,user_ids=singer_ids,only_most_recent=True)
    for qual in singer_quals:
        expiry = QualExpirationData(qual,target_date=date_time)
        part_assignment = GigSongVoicePartAssignment(
                    gig_id=gig.id,
                    user_id=qual.user_id,
                    song_id=qual.song_id,
                    part_id=qual.part_id)
        part_assignment.date_time=qual.date_time
        part_assignment.performance_ready = expiry.valid
        part_assignment.sort_order = expiry.sort_order
        if not expiry.valid:
            part_assignment.sort_order += 1000
        gig.assignments.append(part_assignment)
    assignments = sorted(gig.assignments)

    sorted_songs = sorted([(utility.qual_strength(singer_quals,s),s) for s in songs],reverse=True)
    sort_order = 0
    performable = True
    for ((qual_strength,missing_parts),song) in sorted_songs:
        sort_order += 1
        if qual_strength < 0 and performable:
            performable = False
            setlist_item = GigSetListItem()
            setlist_item.sort_order=sort_order
            setlist_item.comments = "-- END SETLIST --"
            gig.setlist.append(setlist_item)
            sort_order += 100
        if qual_strength < -1:
            break
        setlist_item = GigSetListItem()
        setlist_item.gig_id = gig.id
        setlist_item.song_id = song.id
        setlist_item.duration = song.duration
        setlist_item.song_key = song.performance_key if song.performance_key else song.written_key
        setlist_item.sort_order = sort_order
        gig.setlist.append(setlist_item)
        if performable:
            assign_parts(song,[a for a in assignments if a.song_id==song.id])

    gig.singers = sorted([users_by_id[x] for x in singer_ids])

    session.add(gig)
    # TODO do we need to add the setlist and assignment itens?
    session.commit()
    gigs.append(gig)
    return gig

def reorder_tags(sorted_tag_list):
    for index,tag in enumerate(sorted_tag_list):
        if tag.sort_order != index:
            log.debug(f"reorder_tags: changing sort order for {tag.name} from {tag.sort_order} to {index}")
            tag.sort_order = index
            session.add(tag)
    session.commit()

#def load_from_form_responses(filename,singers=None,songs=None,parts=None,tags=None,caller_session=None):
def load_from_form_responses(filename,singers=None,songs=None,parts=None,tags=None):
    records = []
    today = datetime.datetime.today()
    book = load_workbook(filename, read_only=True)
    for sheetname in book.sheetnames:
        done_with_sheet = False
        log.debug(sheetname)
        sheet = book.get_sheet_by_name(sheetname)
        (labels,*responses) = sheet.rows

        for row in responses:
            if done_with_sheet:
                break
            (timestamp,name,org,song_part_string) = row
            name = name.value
            org = org.value
            song_parts = [x.strip() for x in song_part_string.value.split(",")]
            log.debug(f"{name} - {org} - {song_parts}")
            for song_part in song_parts:
                (song,part,*junk) = song_part.split(" - ")
                if part == 'Baritone':
                    part = 'Bari'
                log.debug(f"{name} {org} - {song} {part}")
                records.append( ('Jamboree', name, org, song, part, today) )
    insert_records(records,singers=singers,songs=songs,parts=parts,tags=tags)

def create_PRE_spreadsheet(filename,singers,songs,tags,parts,quals):
    book = Workbook()
    sheet = book.active
    # remove default sheet
    book.remove_sheet(sheet)
    #new_sheet = book.copy_worksheet(sheet)  # copy from previous sheet
    '''
    for songs to take up more than one cell:
        merge_cells(range_string=None, start_row=None, start_column=None, end_row=None, end_column=None)[source]
        Set merge on a cell range. Range is a cell range (e.g. A1:E1)
    '''
    ###  setup stuff
    SONG_WIDTH = 4
    SONG_SEP_WIDTH = 1
    INITIAL_SONG_OFFSET = 2
    INITIAL_SINGER_OFFSET = 3
    part_names = ['Tenor','Lead','Bari','Bass']
    part_offset = {}
    #for index,part_name in enumerate(part_names):
        #part_offset[part_name] = index
    for part in parts:
        #print(f"PART  {part}")
        part_offset[part.id] = part_names.index(part.name)
    #print(f"part_offset  {part_offset}")
    sorted_singers = []
    singers_by_name = {}
    singer_len = 10
    for singer in singers:
        sorted_singers.append(singer.name)
        singers_by_name[singer.name] = singer
        if len(singer.name) > singer_len:
            singer_len = len(singer.name)
    sorted_singers.sort()
    initial_tag_for_song_id = {}

    for tag in tags:
        #print(f"tag: {tag.name}   id:{tag.id}")
        sheet = book.create_sheet(title=tag.name, index=None)  #index is where to insert sheet
        song_column = {}
        row1 = [None]
        row2 = [None]
        row2_merge_columns = []
        tag_song_list = [(song.id in initial_tag_for_song_id,song.name,song) for song in songs if tag in song.tags]
        tag_song_list.sort()
        for index,song in enumerate([song for (_in_,name,song) in tag_song_list]):
            song_column_temp = INITIAL_SONG_OFFSET + (SONG_WIDTH+SONG_SEP_WIDTH)*index
            if song.id in initial_tag_for_song_id:
                orig_tag = initial_tag_for_song_id[song.id].name
                sub_heading = [f"Qual data in sheet: {orig_tag}",None,None,None]
                row2_merge_columns.append(INITIAL_SONG_OFFSET + (SONG_WIDTH+SONG_SEP_WIDTH)*index)
            else:
                orig_tag = None
                initial_tag_for_song_id[song.id] = tag
                sub_heading = part_names
                song_column[song.id] = song_column_temp
            #print(f" {index}  song: {song.name} {song.id}  column: {song_column_temp}  orig_tag:{orig_tag}")
            row1.extend([song.name,None,None,None]) # extra cells to merge into one big one
            row1.append(None)
            row2.extend(sub_heading)
            row2.append(None)
        sheet.append(row1)
        sheet.append(row2)
        # set colors/styles/sizes for various cells/columns
        sheet.column_dimensions['A'].width = singer_len
        tenor_fill = PatternFill(start_color='FFFFF2CC', end_color='FFFFF2CC', fill_type='solid')
        lead_fill = PatternFill(start_color='FFEFEFEF', end_color='FFEFEFEF', fill_type='solid')
        bari_fill = PatternFill(start_color='FFD9EAD3', end_color='FFD9EAD3', fill_type='solid')
        bass_fill = PatternFill(start_color='FFCFE2F3', end_color='FFCFE2F3', fill_type='solid')
        align_center = Alignment(horizontal='center')
        thin_border = Border(left=Side(style='thin', color='BBBBBB'), 
                right=Side(style='thin', color='BBBBBB'), 
                top=Side(style='thin', color='BBBBBB'),
                bottom=Side(style='thin', color='BBBBBB'))
        for col in song_column.values():
            # song title
            sheet.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col+SONG_WIDTH-1)
            sheet.cell(1,col).alignment = align_center
            # spacer column
            sheet.column_dimensions[get_column_letter(col+4)].width = 2
            # parts columns
            for row in range(2,len(singers)+4):
                cell = sheet.cell(row,col)
                cell.fill = tenor_fill
                cell.border = thin_border
                cell.alignment = align_center
                cell = sheet.cell(row,col+1)
                cell.fill = lead_fill
                cell.border = thin_border
                cell.alignment = align_center
                cell = sheet.cell(row,col+2)
                cell.fill = bari_fill
                cell.border = thin_border
                cell.alignment = align_center
                cell = sheet.cell(row,col+3)
                cell.fill = bass_fill
                cell.border = thin_border
                cell.alignment = align_center
        # take care of songs that have no data in this sheet
        for col in row2_merge_columns:
            sheet.cell(1,col).alignment = align_center
            sheet.cell(2,col).alignment = align_center
            sheet.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col+SONG_WIDTH-1)
            sheet.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col+SONG_WIDTH-1)


        #print(f"song_column:  {song_column}")
        for index,singer_name in enumerate(sorted_singers):
            #print(f"{index} singer: {singer_name}")
            row = INITIAL_SINGER_OFFSET + index
            #sheet.cell(row, column, value=None)  # create/get cell (and set value if not None)
            #put name in first column, regardlesss
            cell = sheet.cell(row, 1, value=singer_name)
            for qual in select_quals(quals, user_ids=[singers_by_name[singer_name].id], only_most_recent=True):
                if qual.song_id in song_column:
                    # if there's no column for a song, ignore since it's not on this sheet
                    col = song_column[qual.song_id] + part_offset[qual.part_id]
                    #print(f"col:{col} DT:{qual.date_time}")
                    cell = sheet.cell(row, col, value=qual.date_time.strftime('%m/%d/%y'))
        #break # for now, just do first tag
                
    book.save(filename)

def load_from_PRE(filename,singers=None,songs=None,parts=None,tags=None):
    # hack to get around both parameter and stuff in loop being named "songs" and "parts"
    db_songs = songs
    db_parts = parts
    records = []
    org = "BTS"
    book = load_workbook(filename, read_only=True)
    for sheetname in book.sheetnames:
        done_with_sheet = False
        #print (sheetname)
        if sheetname in ['Vday Signup', 'Vday quartets', 'Russellville']:
            continue
        sheet = book.get_sheet_by_name(sheetname)
        (songs,parts,*singer_rows) = sheet.rows
        #for part_cell in parts[0:10]:
            #print(f"color: {part_cell.fill.start_color}")
        song_index = []
        part_index = []
        current_song = None
        for song,part in itertools.zip_longest(songs,parts):
            if song is not None and song.value is not None:
                current_song = song.value
            song_index.append(current_song)
            part_index.append(part.value)
        for song_name in song_index:
            log.debug(f"going through songs in sheet: {song_name}")
            if song_name not in [song.name for song in db_songs]:
                if song_name is not None and not song_name.isspace():
                    # create tag and song even if nobody is qual'd
                    log.debug(f"  need to create song: {song_name}")
                    records.append( (sheetname, None, None, song_name, None, None) )

        for row in singer_rows:
            if done_with_sheet:
                break
            for cell_index,cell in enumerate(row):
                if cell_index == 0:
                    name = cell.value
                    if name is None:
                        done_with_sheet = True
                        break
                    continue
                if cell.value is not None:
                    log.debug (name, song_index[cell_index],part_index[cell_index], cell.value)
                    records.append( (sheetname, name, org, song_index[cell_index],part_index[cell_index], cell.value) )
    insert_records(records,singers=singers,songs=db_songs,parts=db_parts,tags=tags)

# here for posterity.  Now this function is handled by 'alembic upgrade head'
def setup_database():
    Base.metadata.create_all(engine)



#engine = create_engine("sqlite://", echo=False)
engine = create_engine("sqlite:///btsbotdb", echo=False)
Session = sessionmaker(engine)
session = Session()
session.expire_on_commit = False

# ugly globals - fixme

users = []
parts = []
songparts = []
songs = []
tags = []
quals = []
gigs = []
songs_by_id = {}
parts_by_id = {}
users_by_id = {}
persistent_values = {}
slack_name_from_db_name = {}
name_map=os.environ.get("SLACK_NAME_MAP")

if name_map:
    for pair in name_map.split(','):
        name_parts = pair.split(':')
        if len(name_parts) != 2:
            log.error(f"SLACK_NAME_MAP env var is expecting pairs of 'db_name:slack_name', comma separated")
            break
        slack_name_from_db_name[name_parts[0]] = name_parts[1]

slack_id_from_slack_name = {}
user_id_from_slack_id = {}

def auth_from_ENV():
    auths = ['admin','music_team','evaluator']
    for auth in auths:
        auth_ENV = f"AUTH_{auth.upper()}"
        auth_name_list=os.environ.get(auth_ENV)
        if auth_name_list:
            for name in auth_name_list.split(','):
                #get user, then set appropriate auth(s)
                try:
                    user = user_from_name(name)
                except:
                    log.error(f"No database user found with name: {name}  (trying to add auth {auth})")
                    continue
                log.info(f"giving user {name} auth {auth}")
                user.add_auth(auth)
        else:
            log.warning(f"{auth_ENV} not specified.  Skipping")

def slack_id_from_user_id(user_id):
    return [x for x in user_id_from_slack_id.keys() if user_id_from_slack_id[x] == user_id][0]

def tag_from_id(tag_id):
    return [x for x in tags if x.id == tag_id][0]

def user_from_name(name):
    users_with_name = [x for x in users if x.name == name]
    if len(users_with_name) == 1:
        return users_with_name[0]
    log.error(f"user_from_name: User not found (or too many) for {name} {users_with_name}")

def set_persistent_value(key,value):
    if key in persistent_values:
        old_kv = persistent_values[key]
        session.delete(old_kv)
    kv = PersistentKeyValue(key,value)
    persistent_values[key] = kv
    session.add(kv)
    session.commit()

def get_persistent_value(key):
    return persistent_values[key].value

def load_database():
    #class PersistentKeyValue(Base):
    #persistent_values = {}
    stmt = select(PersistentKeyValue)
    for kv in session.scalars(stmt):
        persistent_values[kv.key] = kv
    stmt = select(User)
    for user in session.scalars(stmt):
        users.append(user)
        users_by_id[user.id] = user
        if user.slack_id:
            user_id_from_slack_id[user.slack_id] = user.id
    stmt = select(VoicePart)
    for part in session.scalars(stmt):
        parts.append(part)
        parts_by_id[part.id] = part
    stmt = select(SongVoicePart)
    for x in session.scalars(stmt):
        songparts.append(x)
    stmt = select(Song)
    for song in session.scalars(stmt):
        songs.append(song)
        songs_by_id[song.id] = song
    stmt = select(Tag)
    for tag in session.scalars(stmt):
        tags.append(tag)
    stmt = select(SongVoicePartQual)
    for qual in session.scalars(stmt):
        quals.append(qual)
    stmt = select(Gig)
    for gig in session.scalars(stmt):
        gigs.append(gig)

